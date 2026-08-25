"""
Reading the uploaded workbook — **streaming, never wholesale**.

⚠️  `read_only=True` on every single open, with no exception.

    openpyxl's ordinary mode materialises every cell of every sheet as a Python
    object before the first row is handed back. A ten-thousand-row sheet with
    thirty-odd columns is three hundred thousand objects, and three sheets of
    them on a shared host with a hard memory cap is the process being killed
    with no traceback anywhere. Read-only mode yields row tuples instead.

⚠️  And the file is **opened once per chunk and seeked**, not held open between
    requests.

    There is no process that survives from one chunk to the next: the browser
    calls a new request, or the cron runner starts a new interpreter. A file
    handle cannot span that, so each pass re-opens and skips to the cursor.
    Skipping is cheap in read-only mode — it walks the XML without building
    cells — and it is the price of a job that survives a closed laptop.
"""

from __future__ import annotations

import datetime as dt
from collections.abc import Iterator
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from imports import spec

# ═══════════════════════════════════════════════════════════
#  Cell coercion
# ═══════════════════════════════════════════════════════════
#
# ⚠️  Excel hands back types nobody asked for, and every one of them has bitten
#     a real import:
#
#       • a barcode typed as a number arrives as 6.22123e+12 and loses digits
#       • a SKU of "00123" arrives as the integer 123
#       • a price arrives as a float — forbidden by core/money.py in nine words
#       • a date arrives as a datetime, or as the string "30/06/2027",
#         or as the integer 46568 (Excel's own serial)
#
#     So nothing is trusted: every cell is coerced by **the column's declared
#     kind**, not by what Excel claims the cell is.

#: Values that mean "yes" in a sheet filled by an Arabic-speaking human.
TRUE_WORDS = frozenset(
    {"1", "true", "yes", "y", "نعم", "صح", "متاح", "مفعل", "مفعّل", "✓", "x"}
)
FALSE_WORDS = frozenset({"0", "false", "no", "n", "لا", "خطأ", "غير مفعل", "غير مفعّل", "-"})

#: Arabic-Indic and Eastern Arabic-Indic digits → ASCII.
#
# ⚠️  A price typed on an Arabic keyboard arrives as "٤٥٫٠٠" and every numeric
#     parser in Python rejects it. Rejecting the row instead of translating the
#     digits makes the importer unusable for exactly the people it is for.
_DIGIT_MAP = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹٫٬", "01234567890123456789.,")


class CellError(ValueError):
    """A value that cannot be coerced. Carries a message the admin can act on."""


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        # ⚠️  A SKU or barcode Excel decided was a number. `str(6221234567890.0)`
        #     keeps the '.0'; int() first is what preserves the identifier.
        return str(int(value))
    if isinstance(value, dt.datetime | dt.date):
        return value.isoformat()
    return str(value).strip()


def _normalise_digits(raw: str) -> str:
    return raw.translate(_DIGIT_MAP).replace(",", "").replace(" ", "").replace("‏", "")


def coerce(value, kind: str, *, column_header: str = "") -> object:
    """Turn one raw cell into the type the column declares, or raise `CellError`."""
    if kind == spec.Kind.BOOL:
        raw = _text(value).lower()
        if raw == "":
            return None
        if raw in TRUE_WORDS:
            return True
        if raw in FALSE_WORDS:
            return False
        raise CellError(f"«{raw}» ليست نعم أو لا")

    if kind in (spec.Kind.INT, spec.Kind.MONEY):
        raw = _normalise_digits(_text(value))
        if raw == "":
            return None
        try:
            # ⚠️  `Decimal` even for the integer case, then converted.
            #     `int("12.0")` raises while `int(Decimal("12.0"))` is 12 — and
            #     "12.0" is exactly what Excel writes for a whole number.
            number = Decimal(raw)
        except (InvalidOperation, ArithmeticError) as exc:
            raise CellError(f"«{raw}» ليست قيمة رقمية") from exc

        if kind == spec.Kind.INT:
            if number != number.to_integral_value():
                raise CellError(f"«{raw}» يجب أن يكون رقمًا صحيحًا")
            return int(number)

        # ⚠️  Quantised here, not left to the database.
        #     A price of 45.006 is silently rounded on save, so the admin's file
        #     and the stored catalogue disagree with nothing to explain why.
        return number.quantize(Decimal("0.01"))

    if kind == spec.Kind.DATE:
        return _coerce_date(value)

    if kind == spec.Kind.ATTRS:
        return _coerce_attributes(value)

    return _text(value)


def _coerce_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value

    raw = _normalise_digits(_text(value)).replace("/", "-").replace(".", "-")
    if not raw:
        return None

    for pattern in ("%Y-%m-%d", "%d-%m-%Y", "%m-%d-%Y", "%Y-%m", "%d-%m-%y"):
        try:
            return dt.datetime.strptime(raw, pattern).date()
        except ValueError:
            continue

    raise CellError(f"«{_text(value)}» ليست تاريخًا — الصيغة المتوقعة YYYY-MM-DD")


def _coerce_attributes(value) -> dict:
    """`size=M;color=blue` → `{'size': 'M', 'color': 'blue'}`"""
    raw = _text(value)
    if not raw:
        return {}

    attributes: dict[str, str] = {}
    for pair in raw.replace("،", ";").replace(",", ";").split(";"):
        pair = pair.strip()
        if not pair:
            continue
        if "=" not in pair:
            raise CellError(f"«{pair}» ليست زوجًا بصيغة مفتاح=قيمة")
        name, _, item = pair.partition("=")
        name = name.strip()
        if not name:
            raise CellError("اسم الخاصية فارغ")
        attributes[name] = item.strip()
    return attributes


# ═══════════════════════════════════════════════════════════
#  Header matching
# ═══════════════════════════════════════════════════════════


def _normalise_header(raw: str) -> str:
    """
    ⚠️  A header survives a round trip through Google Sheets, a copy-paste and a
        well-meaning admin who removed the asterisk. Matching on the exact
        string meant a file that looked identical failed with "column missing",
        and nobody could see the difference on screen.
    """
    text = _text(raw).lower()
    for noise in ("*", "‏", "‎", " "):
        text = text.replace(noise, " ")
    return " ".join(text.split())


def _header_index(header_cells, key_cells, columns) -> dict[str, int]:
    """
    Map `column.key` → position, preferring the hidden key row.

    ⚠️  Two rows are consulted in order, and the order is the point: the English
        key row is authoritative because it is stable, and the Arabic header row
        is the fallback for a file whose hidden row was deleted. Trusting the
        Arabic row first means a translated header silently binds to the wrong
        column — which writes prices into weights.
    """
    index: dict[str, int] = {}

    by_key = {_normalise_header(cell): position for position, cell in enumerate(key_cells)}
    by_header = {_normalise_header(cell): position for position, cell in enumerate(header_cells)}

    for column in columns:
        position = by_key.get(column.key.lower())
        if position is None:
            position = by_header.get(_normalise_header(column.header))
        if position is not None:
            index[column.key] = position

    return index


# ═══════════════════════════════════════════════════════════
#  The workbook
# ═══════════════════════════════════════════════════════════


@dataclass
class SheetShape:
    """What a sheet turned out to contain."""

    name: str
    present: bool
    row_count: int = 0
    index: dict[str, int] | None = None
    missing_required: list[str] = None  # type: ignore[assignment]

    def __post_init__(self):
        if self.index is None:
            self.index = {}
        if self.missing_required is None:
            self.missing_required = []


@dataclass
class Row:
    """One data row, keyed by column key, with the Excel row number attached."""

    number: int
    values: dict[str, object]
    #: Cells that failed coercion — reported without stopping the row.
    cell_errors: list[tuple[str, str, str]]  # (key, raw, message)

    def get(self, key: str, default=None):
        value = self.values.get(key)
        return default if value in (None, "") else value

    @property
    def identifier(self) -> str:
        return str(self.values.get("sku") or self.values.get("parent_sku") or "")


def inspect(path_or_file, required: dict[str, list[str]] | None = None) -> dict[str, SheetShape]:
    """
    Open the workbook, confirm its shape, and count rows — **without reading values**.

    ⚠️  This runs at upload, before anything is stored as a job worth resuming.
        A file missing the `products` sheet, or missing `sku`, is rejected in one
        pass with a message naming what is absent — rather than after the admin
        has waited through a dry run of ten thousand rows.

    ⚠️  `required` comes from the caller because it depends on the import mode,
        which this file knows nothing about. See `spec.required_keys_for`: a
        price-list update legitimately carries two columns.
    """
    required = required or {name: spec.required_keys(name) for name in spec.SHEETS}
    workbook = load_workbook(path_or_file, read_only=True, data_only=True)
    try:
        shapes: dict[str, SheetShape] = {}
        titles = set(workbook.sheetnames)

        for sheet_name, columns in spec.SHEETS.items():
            if sheet_name not in titles:
                shapes[sheet_name] = SheetShape(name=sheet_name, present=False)
                continue

            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(values_only=True)

            header_cells = next(rows, ()) or ()
            key_cells = next(rows, ()) or ()
            index = _header_index(header_cells, key_cells, columns)

            missing = [
                key for key in required.get(sheet_name, ()) if key not in index
            ]

            # ⚠️  Counted by walking, not by `worksheet.max_row`.
            #
            #     In read-only mode `max_row` comes from the file's own dimension
            #     record, which Excel writes optimistically and other tools omit
            #     entirely: it reported 1,048,576 for a sheet holding nine
            #     hundred rows, and the progress bar was a flat line at 0%.
            count = sum(1 for row in rows if any(cell not in (None, "") for cell in row))

            shapes[sheet_name] = SheetShape(
                name=sheet_name,
                present=True,
                row_count=count,
                index=index,
                missing_required=missing,
            )

        return shapes
    finally:
        # ⚠️  Read-only workbooks hold the zip handle open. On Windows an
        #     unclosed handle means the uploaded file cannot be deleted or
        #     replaced, and the failure surfaces much later as a permission error.
        workbook.close()


def read_rows(
    path_or_file,
    sheet_name: str,
    *,
    start: int = 0,
    limit: int | None = None,
    only: frozenset[str] | None = None,
) -> Iterator[Row]:
    """
    Yield coerced rows from one sheet.

    `start` is a count of data rows already handled — the job's cursor — and
    `limit` is the chunk size. Together they make one pass resumable.

    ⚠️  `only` narrows the work to a few columns, and it exists for the scans
        that need one.

        Answering "which SKUs does this file contain?" over fifty thousand rows
        should not coerce thirty other columns per row — the parse still walks
        the XML, but the coercion is where the time actually goes.
    """
    columns = spec.columns_for(sheet_name)
    workbook = load_workbook(path_or_file, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return

        worksheet = workbook[sheet_name]
        raw_rows = worksheet.iter_rows(values_only=True)

        header_cells = next(raw_rows, ()) or ()
        key_cells = next(raw_rows, ()) or ()
        index = _header_index(header_cells, key_cells, columns)
        if only is not None:
            index = {key: position for key, position in index.items() if key in only}
        by_key = spec.column_map(sheet_name)

        seen = 0
        yielded = 0

        for offset, cells in enumerate(raw_rows):
            # ⚠️  A blank row is skipped and **not counted**.
            #
            #     Spreadsheets are full of them — a gap between suppliers, a
            #     trailing region Excel keeps. Counting them makes the cursor and
            #     the row total disagree, and the job stops one chunk short of
            #     the end and reports itself finished.
            if not cells or all(cell in (None, "") for cell in cells):
                continue

            seen += 1
            if seen <= start:
                continue

            values: dict[str, object] = {}
            cell_errors: list[tuple[str, str, str]] = []

            for key, position in index.items():
                raw = cells[position] if position < len(cells) else None
                column = by_key[key]
                try:
                    values[key] = coerce(raw, column.kind, column_header=column.header)
                except CellError as exc:
                    values[key] = None
                    cell_errors.append((key, _text(raw)[:200], str(exc)))

            yield Row(
                # ⚠️  `spec.FIRST_DATA_ROW + offset`, not a counter.
                #
                #     The number reported must be the number the admin sees in
                #     the Excel row gutter, blank rows included — a report that
                #     says "row 812" about row 840 is worse than no report.
                number=spec.FIRST_DATA_ROW + offset,
                values=values,
                cell_errors=cell_errors,
            )

            yielded += 1
            if limit is not None and yielded >= limit:
                return
    finally:
        workbook.close()
