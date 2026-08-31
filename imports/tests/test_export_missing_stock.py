"""
Exporting the products that have no opening balance.

⚠️  The file this command writes is fed straight back into the importer, so the
    test drives it through the **real** pipeline — upload · dry run · execute —
    rather than asserting on cell values.

    A spreadsheet that looks right and is read wrong is exactly the failure the
    command already had once: a single header row instead of the template's two
    made the reader treat the first product as the key row and drop it, silently,
    with every remaining row perfectly correct.
"""

import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from openpyxl import load_workbook

from catalog.models import Category, Product
from imports import services, spec
from imports.models import ImportJob, ImportStatus
from inventory.models import Stock, StockLocation

pytestmark = pytest.mark.django_db


def export(tmp_path, **options) -> bytes:
    destination = tmp_path / "opening-stock.xlsx"
    call_command("export_missing_stock", out=str(destination), verbosity=0, **options)
    return destination.read_bytes()


def run(job: ImportJob) -> ImportJob:
    for _ in range(400):
        if not job.is_running:
            return job
        services.advance(job)
    raise AssertionError(f"job never stopped — {job.status}/{job.phase}")


@pytest.fixture
def location(db):
    return StockLocation.objects.create(
        code="main", name_ar="الرئيسي", name_en="Main", is_default=True, is_sellable=True
    )


@pytest.fixture
def category(db):
    return Category.objects.create(name_ar="فئة", name_en="Category")


def make_product(category, sku, **overrides):
    return Product.objects.create(
        sku=sku,
        name_ar=f"صنف {sku}",
        name_en=f"Item {sku}",
        category=category,
        base_price=Decimal("10.00"),
        **overrides,
    )


class TestExport:
    def test_it_lists_exactly_the_products_with_no_stock_row(self, tmp_path, category, location):
        stocked = make_product(category, "HAS-1")
        make_product(category, "NONE-1")
        make_product(category, "NONE-2")
        Stock.objects.create(product=stocked, location=location, quantity_physical=5)

        worksheet = load_workbook(io.BytesIO(export(tmp_path)))[spec.Sheet.STOCK]
        skus = {
            row[0]
            for row in worksheet.iter_rows(min_row=spec.FIRST_DATA_ROW, max_col=1, values_only=True)
            if row[0]
        }

        assert skus == {"NONE-1", "NONE-2"}

    def test_a_row_that_reads_zero_is_not_a_missing_balance(self, tmp_path, category, location):
        """
        ⚠️  Received and then sold out is a different fact from never received.

            Both are invisible to a customer; only the second belongs in an
            opening-balance file. Putting the first there invites a second
            "opening" balance for a product that already has a ledger.
        """
        sold_out = make_product(category, "SOLD-1")
        Stock.objects.create(product=sold_out, location=location, quantity_physical=0)
        # A genuinely un-received one, so there is a file to read at all
        make_product(category, "NEVER-1")

        worksheet = load_workbook(io.BytesIO(export(tmp_path)))[spec.Sheet.STOCK]
        skus = {
            row[0]
            for row in worksheet.iter_rows(min_row=spec.FIRST_DATA_ROW, max_col=1, values_only=True)
            if row[0]
        }

        assert skus == {"NEVER-1"}

    def test_inactive_products_are_excluded_unless_asked_for(self, tmp_path, category):
        """
        ⚠️  A discontinued product needs no opening balance.

            Including it by default puts rows in front of the admin for items
            nobody intends to sell, and the file's length is the only measure
            they have of how much work is left.
        """
        make_product(category, "OFF-1", is_active=False)

        # Nothing active is missing a balance, so the default export writes no file
        assert not (tmp_path / "opening-stock.xlsx").exists()
        call_command("export_missing_stock", out=str(tmp_path / "opening-stock.xlsx"), verbosity=0)
        assert not (tmp_path / "opening-stock.xlsx").exists()

        included = load_workbook(io.BytesIO(export(tmp_path, include_inactive=True)))[
            spec.Sheet.STOCK
        ]
        assert included.cell(row=spec.FIRST_DATA_ROW, column=1).value == "OFF-1"

    def test_an_unsellable_location_is_refused(self, tmp_path, category, db):
        """
        ⚠️  A balance received into quarantine leaves the product just as invisible
            — which is the whole problem this file exists to undo.
        """
        from django.core.management.base import CommandError

        StockLocation.objects.create(
            code="quarantine", name_ar="حجر", name_en="Quarantine", is_sellable=False
        )
        make_product(category, "P-1")

        with pytest.raises(CommandError, match="غير قابل للبيع"):
            export(tmp_path, location="quarantine")


class TestRoundTrip:
    """The file goes back into the importer with no transformation."""

    def test_the_importer_accepts_it_and_receives_the_balances(self, tmp_path, category, location):
        """
        ⚠️  **The end-to-end guarantee**: export ← fill in ← import ← visible.

            And the count is asserted product by product, because the failure this
            guards against loses exactly one row and leaves the rest right.
        """
        for index in range(3):
            make_product(category, f"NEW-{index}")

        payload = export(tmp_path, location=location.code)

        # The admin fills in the two highlighted columns
        book = load_workbook(io.BytesIO(payload))
        worksheet = book[spec.Sheet.STOCK]
        keys = [cell.value for cell in worksheet[spec.KEY_ROW]]
        quantity_column = keys.index("quantity") + 1
        cost_column = keys.index("unit_cost") + 1

        for row in range(spec.FIRST_DATA_ROW, spec.FIRST_DATA_ROW + 3):
            worksheet.cell(row=row, column=quantity_column, value=7)
            worksheet.cell(row=row, column=cost_column, value=Decimal("4.50"))

        buffer = io.BytesIO()
        book.save(buffer)

        job = services.create_job(
            SimpleUploadedFile("opening-stock.xlsx", buffer.getvalue()),
            mode=spec.ImportMode.UPDATE_ONLY,
        )

        services.start_validation(job)
        job = run(job)
        assert job.status == ImportStatus.VALIDATED, job.error_message

        services.start_execution(job)
        job = run(job)
        assert job.status == ImportStatus.DONE, job.error_message

        # ⚠️  Every one of the three — not "some stock was created"
        for index in range(3):
            product = Product.objects.get(sku=f"NEW-{index}")
            row = Stock.objects.get(product=product, location=location)
            assert row.quantity_physical == 7, f"NEW-{index}"

    def test_an_empty_export_writes_nothing(self, tmp_path, category, location, capsys):
        stocked = make_product(category, "HAS-1")
        Stock.objects.create(product=stocked, location=location, quantity_physical=5)

        destination = tmp_path / "opening-stock.xlsx"
        call_command("export_missing_stock", out=str(destination))

        # ⚠️  No file at all rather than an empty one — a workbook with only
        #     headers looks like an export that lost its rows.
        assert not destination.exists()
        assert "كل المنتجات لها رصيد" in capsys.readouterr().out
