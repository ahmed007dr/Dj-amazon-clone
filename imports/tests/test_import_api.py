"""
The endpoints: permissions, ownership, and the shape the frontend depends on.

⚠️  The permission tests here are the point of the file. A merchandiser who may
    add products must not receive a million pounds of stock because it rode
    along in the same upload — and that is a rule about the **file**, not about
    the endpoint, which is why it can only be tested through one.
"""

import pytest
from django.apps import apps
from django.urls import reverse
from rest_framework.test import APIClient

from accounts.models import AccountType, User
from catalog.models import Product
from imports import services, spec
from imports.models import ImportStatus
from imports.tests.conftest import PASSWORD, product_row, workbook

pytestmark = pytest.mark.django_db

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def post_file(client, payload: bytes, *, mode=spec.ImportMode.CREATE_ONLY, name="c.xlsx"):
    from django.core.files.uploadedfile import SimpleUploadedFile

    return client.post(
        reverse("v1:imports:jobs"),
        {"file": SimpleUploadedFile(name, payload, content_type=XLSX), "mode": mode},
        format="multipart",
    )


def drive(client, job_id: str) -> dict:
    """The loop the browser runs — `is_running`, not `not is_terminal`."""
    url = reverse("v1:imports:job-advance", args=[job_id])
    for _ in range(400):
        body = client.post(url).json()
        if not body["is_running"]:
            return body
    raise AssertionError("job never stopped")


# ═══════════════════════════════════════════════════════════
#  The template and the guide
# ═══════════════════════════════════════════════════════════


class TestTemplate:
    def test_the_template_downloads_as_a_workbook(self, admin_client, category):
        response = admin_client.get(reverse("v1:imports:template"))

        assert response.status_code == 200
        assert response["Content-Type"] == XLSX
        assert "attachment" in response["Content-Disposition"]
        # ⚠️  A real zip signature — a template that opens with a repair prompt
        #     is worse than no template.
        assert response.content[:2] == b"PK"

    def test_the_template_is_never_cached(self, admin_client, category):
        """It is built from the live catalogue; a proxy copy is another admin's."""
        response = admin_client.get(reverse("v1:imports:template"))
        assert response["Cache-Control"] == "no-store"

    def test_the_spec_describes_the_columns_the_template_carries(self, admin_client, category):
        body = admin_client.get(reverse("v1:imports:spec")).json()

        sheets = {sheet["name"]: sheet for sheet in body["sheets"]}
        assert sheets["products"]["required"] is True
        assert sheets["stock"]["required"] is False

        keys = [column["key"] for column in sheets["products"]["columns"]]
        assert keys == [column.key for column in spec.PRODUCT_COLUMNS]

    def test_the_spec_flags_a_store_that_cannot_import_yet(self, admin_client, db):
        """
        ⚠️  "Why did every single row fail?" has this answer more often than any
            other, and the screen can say so before the admin uploads anything.
        """
        body = admin_client.get(reverse("v1:imports:spec")).json()
        assert body["blocking"]["no_categories"] is True

    def test_an_anonymous_visitor_gets_nothing(self, db):
        assert APIClient().get(reverse("v1:imports:template")).status_code in (401, 403)


# ═══════════════════════════════════════════════════════════
#  Upload
# ═══════════════════════════════════════════════════════════


class TestUpload:
    def test_upload_reports_the_row_counts(self, admin_client, category):
        payload = workbook({"products": [product_row(f"S{n}", category) for n in range(3)]})
        response = post_file(admin_client, payload)

        assert response.status_code == 201
        body = response.json()
        assert body["row_counts"] == {"products": 3}
        assert body["status"] == ImportStatus.UPLOADED
        assert body["duplicate_of"] is None

    def test_a_non_xlsx_file_is_refused_by_name_before_it_reaches_a_parser(
        self, admin_client, category
    ):
        """
        ⚠️  An extension is a claim the client makes, so the content is checked
            too — but rejecting on the name first costs one comparison and keeps
            a renamed executable away from a zip parser.
        """
        response = post_file(admin_client, b"not a workbook", name="payload.exe")
        assert response.status_code == 400

    def test_a_file_named_xlsx_that_is_not_one_is_refused_by_content(
        self, admin_client, category
    ):
        response = post_file(admin_client, b"not a workbook at all", name="c.xlsx")
        assert response.status_code == 400
        assert response.json()["code"] == "IMPORT_FILE_UNREADABLE"

    def test_re_uploading_the_same_bytes_is_flagged(self, admin_client, category):
        payload = workbook({"products": [product_row("A", category)]})

        first = post_file(admin_client, payload).json()
        job = services.ImportJob.objects.get(pk=first["id"])
        services.start_validation(job)
        for _ in range(20):
            if not job.is_running:
                break
            services.advance(job)
        services.start_execution(job)
        for _ in range(20):
            if not job.is_running:
                break
            services.advance(job)

        second = post_file(admin_client, payload).json()
        assert second["duplicate_of"]["id"] == first["id"]


# ═══════════════════════════════════════════════════════════
#  The lifecycle over HTTP
# ═══════════════════════════════════════════════════════════


class TestLifecycle:
    def test_the_full_loop(self, admin_client, category):
        payload = workbook({"products": [product_row(f"S{n}", category) for n in range(4)]})
        job_id = post_file(admin_client, payload).json()["id"]

        admin_client.post(reverse("v1:imports:job-validate", args=[job_id]))
        body = drive(admin_client, job_id)
        assert body["status"] == ImportStatus.VALIDATED
        assert body["preview"]["summary"]["will_create"] == 4
        assert body["progress_percent"] == 100
        assert Product.objects.count() == 0

        admin_client.post(reverse("v1:imports:job-execute", args=[job_id]))
        body = drive(admin_client, job_id)
        assert body["status"] == ImportStatus.DONE
        assert Product.objects.count() == 4

        published = admin_client.post(reverse("v1:imports:job-publish", args=[job_id])).json()
        assert published["published"] == 4

    def test_execute_before_validate_is_refused(self, admin_client, category):
        payload = workbook({"products": [product_row("A", category)]})
        job_id = post_file(admin_client, payload).json()["id"]

        response = admin_client.post(reverse("v1:imports:job-execute", args=[job_id]))
        assert response.status_code == 400
        assert response.json()["code"] == "IMPORT_NOT_VALIDATED"

    def test_advancing_a_finished_job_is_harmless(self, admin_client, category):
        """
        ⚠️  The browser loop and the periodic runner can both call this on the
            same job. A second call has to be a no-op, not a second import.
        """
        payload = workbook({"products": [product_row("A", category)]})
        job_id = post_file(admin_client, payload).json()["id"]
        admin_client.post(reverse("v1:imports:job-validate", args=[job_id]))
        drive(admin_client, job_id)
        admin_client.post(reverse("v1:imports:job-execute", args=[job_id]))
        drive(admin_client, job_id)

        assert Product.objects.count() == 1
        admin_client.post(reverse("v1:imports:job-advance", args=[job_id]))
        assert Product.objects.count() == 1

    def test_another_admin_may_not_drive_someone_elses_job(self, admin_client, category, db):
        """
        ⚠️  Two browsers on one job step on the same cursor: both read 4,000,
            both process the same rows, and one chunk imports twice. Ownership
            is the cheapest possible lock.
        """
        from core.testing import grant_all_domains

        payload = workbook({"products": [product_row("A", category)]})
        job_id = post_file(admin_client, payload).json()["id"]

        other = User.objects.create_user(
            email="other@test.local", password=PASSWORD, account_type=AccountType.ADMIN
        )
        other.is_active = True
        other.save()
        apps.get_model("administration", "AdminProfile").objects.create(user=other)
        grant_all_domains(other)

        client = APIClient()
        client.force_authenticate(user=other)

        response = client.post(reverse("v1:imports:job-validate", args=[job_id]))
        assert response.status_code == 403

    def test_cancelling_stops_the_job(self, admin_client, category):
        payload = workbook({"products": [product_row(f"S{n}", category) for n in range(3)]})
        job_id = post_file(admin_client, payload).json()["id"]
        admin_client.post(reverse("v1:imports:job-validate", args=[job_id]))
        drive(admin_client, job_id)

        body = admin_client.post(reverse("v1:imports:job-cancel", args=[job_id])).json()
        assert body["status"] == ImportStatus.CANCELLED
        assert body["is_running"] is False


# ═══════════════════════════════════════════════════════════
#  The error report
# ═══════════════════════════════════════════════════════════


class TestErrors:
    @pytest.fixture
    def rejected(self, admin_client, category):
        payload = workbook(
            {
                "products": [
                    product_row("GOOD", category),
                    product_row("BAD", category, category_path="no/such/path"),
                ]
            }
        )
        job_id = post_file(admin_client, payload).json()["id"]
        admin_client.post(reverse("v1:imports:job-validate", args=[job_id]))
        drive(admin_client, job_id)
        return job_id

    def test_errors_are_paginated_with_the_row_the_admin_sees(self, admin_client, rejected):
        body = admin_client.get(reverse("v1:imports:job-errors", args=[rejected])).json()

        assert body["count"] == 1
        error = body["results"][0]
        assert error["identifier"] == "BAD"
        assert error["row_number"] == spec.FIRST_DATA_ROW + 1

    def test_the_error_file_is_a_workbook_the_admin_can_re_upload(
        self, admin_client, rejected, category
    ):
        """
        ⚠️  The whole point: the failing rows come back as a valid import file,
            so the loop is fix-in-place and upload — not copy corrections back
            into a ten-thousand-row original by hand.
        """
        response = admin_client.get(reverse("v1:imports:job-error-file", args=[rejected]))

        assert response.status_code == 200
        assert response["Content-Type"] == XLSX
        assert response.content[:2] == b"PK"

        import io

        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(response.content))
        sheet = book["products"]
        # The hidden key row survives, which is what makes it re-uploadable.
        assert sheet.cell(row=spec.KEY_ROW, column=1).value == "sku"
        # Exactly the failing row, with the value the admin actually typed.
        assert sheet.cell(row=spec.FIRST_DATA_ROW, column=1).value == "BAD"
        assert sheet.cell(row=spec.FIRST_DATA_ROW, column=4).value == "no/such/path"
        book.close()

    def test_a_clean_job_has_no_error_file(self, admin_client, category):
        payload = workbook({"products": [product_row("A", category)]})
        job_id = post_file(admin_client, payload).json()["id"]
        admin_client.post(reverse("v1:imports:job-validate", args=[job_id]))
        drive(admin_client, job_id)

        response = admin_client.get(reverse("v1:imports:job-error-file", args=[job_id]))
        assert response.status_code == 404


# ═══════════════════════════════════════════════════════════
#  Permissions
# ═══════════════════════════════════════════════════════════


class TestPermissions:
    @pytest.fixture
    def catalog_only(self, db):
        """An admin who may manage the catalogue and nothing else."""
        from django.contrib.auth.models import Permission

        user = User.objects.create_user(
            email="merch@test.local", password=PASSWORD, account_type=AccountType.ADMIN
        )
        user.is_active = True
        user.save()
        apps.get_model("administration", "AdminProfile").objects.create(user=user)
        permission = Permission.objects.get(
            content_type__app_label="catalog", codename="change_product"
        )
        user.user_permissions.add(permission)

        client = APIClient()
        client.force_authenticate(user=user)
        return client

    def test_a_products_only_file_needs_only_the_catalogue_permission(
        self, catalog_only, category
    ):
        payload = workbook({"products": [product_row("A", category)]})
        assert post_file(catalog_only, payload).status_code == 201

    def test_a_file_carrying_stock_needs_the_inventory_permission_too(
        self, catalog_only, category, location
    ):
        """
        ⚠️  Opening stock writes batches, costs and ledger movements. Letting it
            in on the catalogue permission alone means a merchandiser receiving
            a million pounds of goods in a column nobody reviewed.
        """
        payload = workbook(
            {
                "products": [product_row("A", category)],
                "stock": [{"sku": "A", "quantity": 5, "unit_cost": "1.00"}],
            }
        )
        response = post_file(catalog_only, payload)

        assert response.status_code == 403
        assert "المخزون" in response.json()["detail"]

    def test_the_rejected_upload_leaves_no_job_behind(self, catalog_only, category, location):
        from imports.models import ImportJob

        payload = workbook(
            {
                "products": [product_row("A", category)],
                "stock": [{"sku": "A", "quantity": 5, "unit_cost": "1.00"}],
            }
        )
        post_file(catalog_only, payload)
        assert ImportJob.objects.filter(deleted_at__isnull=True).count() == 0

    def test_an_admin_with_both_permissions_may_upload_stock(
        self, admin_client, category, location
    ):
        payload = workbook(
            {
                "products": [product_row("A", category)],
                "stock": [{"sku": "A", "quantity": 5, "unit_cost": "1.00"}],
            }
        )
        assert post_file(admin_client, payload).status_code == 201
