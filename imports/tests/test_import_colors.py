"""
Colours on the variants sheet.

⚠️  The colour is described **on each of its variants** (name, hex, price), found or
    created per product by name — so the same colour typed on three sizes stays one
    colour, a re-import updates it rather than minting a second, and a blank cell
    means "leave it" (the sheet has no way to say "clear").
"""

from decimal import Decimal

import pytest

from catalog.models import ProductColor, ProductVariant
from imports import spec
from imports.models import ImportStatus
from imports.tests.conftest import workbook
from imports.tests.test_import_flow import executed, upload, validated

pytestmark = pytest.mark.django_db


def variant_row(sku, parent="GLV", **extra):
    row = {
        "parent_sku": parent,
        "sku": sku,
        "name_ar": f"نسخة {sku}",
        "name_en": f"Variant {sku}",
        "color_name_ar": "أزرق",
        "color_name_en": "Blue",
        "color_hex": "#1E40AF",
    }
    row.update(extra)
    return row


def run(make_product_row, variants, *, products=None, mode=spec.ImportMode.CREATE_ONLY):
    payload = workbook({"products": products or [make_product_row("GLV")], "variants": variants})
    return executed(validated(upload(payload, mode=mode)))


def rejected(make_product_row, variants, **kwargs):
    payload = workbook(
        {"products": kwargs.get("products") or [make_product_row("GLV")], "variants": variants}
    )
    return validated(upload(payload))


class TestCreate:
    def test_a_colour_is_created_and_the_variants_are_linked(self, make_product_row):
        job = run(
            make_product_row,
            [variant_row("GLV-M", attributes="size=M"), variant_row("GLV-L", attributes="size=L")],
        )

        assert job.status == ImportStatus.DONE
        color = ProductColor.objects.get(product__sku="GLV")
        assert (color.name_ar, color.name_en, color.hex) == ("أزرق", "Blue", "#1E40AF")
        assert ProductColor.objects.count() == 1, "نفس اللون على مقاسين = لون واحد"
        assert set(color.variants.values_list("sku", flat=True)) == {"GLV-M", "GLV-L"}

    def test_the_colours_own_price_is_stored(self, make_product_row):
        run(make_product_row, [variant_row("GLV-M", price_override="120.00")])

        assert ProductColor.objects.get().price_override == Decimal("120.00")

    def test_a_dry_run_writes_nothing(self, make_product_row):
        job = rejected(make_product_row, [variant_row("GLV-M")])

        assert job.status == ImportStatus.VALIDATED
        assert not ProductColor.objects.exists() and not ProductVariant.objects.exists()

    def test_a_variant_without_colour_cells_has_no_colour(self, make_product_row):
        run(
            make_product_row,
            [
                {
                    "parent_sku": "GLV",
                    "sku": "GLV-XL",
                    "name_ar": "كبير جدًا",
                    "name_en": "XL",
                    "attributes": "size=XL",
                }
            ],
        )

        assert ProductVariant.objects.get(sku="GLV-XL").color_id is None
        assert not ProductColor.objects.exists()

    def test_one_language_is_copied_to_the_other(self, make_product_row):
        run(make_product_row, [variant_row("GLV-M", color_name_en=None)])

        color = ProductColor.objects.get()
        assert color.name_en == "أزرق" and color.name_ar == "أزرق"

    def test_a_hex_without_the_hash_is_accepted(self, make_product_row):
        run(make_product_row, [variant_row("GLV-M", color_hex="1e40af")])

        assert ProductColor.objects.get().hex == "#1E40AF"

    def test_no_hex_gets_the_neutral_grey(self, make_product_row):
        run(make_product_row, [variant_row("GLV-M", color_hex=None)])

        assert ProductColor.objects.get().hex == "#9CA3AF"

    def test_two_products_may_each_have_a_blue(self, make_product_row):
        run(
            make_product_row,
            [variant_row("A-M", parent="A"), variant_row("B-M", parent="B")],
            products=[make_product_row("A"), make_product_row("B")],
        )

        assert ProductColor.objects.count() == 2


class TestRejected:
    @pytest.mark.parametrize("bad", ["blue", "#12345", "#GGGGGG", "#1234567"])
    def test_a_malformed_hex(self, make_product_row, bad):
        job = rejected(make_product_row, [variant_row("GLV-M", color_hex=bad)])

        assert job.status == ImportStatus.REJECTED
        assert job.errors.filter(message__contains="#RRGGBB").exists()

    def test_a_hex_with_no_name(self, make_product_row):
        job = rejected(
            make_product_row,
            [variant_row("GLV-M", color_name_ar=None, color_name_en=None)],
        )

        assert job.status == ImportStatus.REJECTED
        assert job.errors.filter(message__contains="بلا اسم").exists()

    def test_a_negative_price(self, make_product_row):
        job = rejected(make_product_row, [variant_row("GLV-M", price_override="-5.00")])

        assert job.status == ImportStatus.REJECTED

    def test_one_colour_described_two_ways(self, make_product_row):
        job = rejected(
            make_product_row,
            [variant_row("GLV-M", color_hex="#111111"), variant_row("GLV-L", color_hex="#222222")],
        )

        assert job.status == ImportStatus.REJECTED
        assert job.errors.filter(message__contains="يخالف").exists()
        assert not ProductColor.objects.exists()


class TestReimport:
    def test_a_second_import_reuses_the_colour_and_updates_what_it_is_given(self, make_product_row):
        run(make_product_row, [variant_row("GLV-M", price_override="120.00")])

        job = run(
            make_product_row,
            [variant_row("GLV-M", color_hex="#000080", price_override="130.00")],
            mode=spec.ImportMode.UPSERT,
        )

        assert job.status == ImportStatus.DONE
        color = ProductColor.objects.get()
        assert (color.hex, color.price_override) == ("#000080", Decimal("130.00"))
        assert ProductVariant.objects.get(sku="GLV-M").color_id == color.pk

    def test_a_blank_cell_leaves_the_colour_alone(self, make_product_row):
        run(make_product_row, [variant_row("GLV-M", price_override="120.00")])

        run(
            make_product_row,
            [variant_row("GLV-M", color_hex=None, price_override=None)],
            mode=spec.ImportMode.UPSERT,
        )

        color = ProductColor.objects.get()
        assert (color.hex, color.price_override) == ("#1E40AF", Decimal("120.00"))

    def test_a_new_size_joins_the_existing_colour(self, make_product_row):
        run(make_product_row, [variant_row("GLV-M")])

        run(make_product_row, [variant_row("GLV-L")], mode=spec.ImportMode.UPSERT)

        assert ProductColor.objects.count() == 1
        assert ProductColor.objects.get().variants.count() == 2

    def test_a_file_without_the_colour_columns_does_not_touch_colours(self, make_product_row):
        run(make_product_row, [variant_row("GLV-M")])
        old_columns = ["parent_sku", "sku", "name_ar", "name_en", "price_adjustment"]
        payload = workbook(
            {
                "products": [make_product_row("GLV")],
                "variants": [
                    {
                        "parent_sku": "GLV",
                        "sku": "GLV-M",
                        "name_ar": "جديد",
                        "name_en": "New",
                        "price_adjustment": "3.00",
                    }
                ],
            },
            columns={"variants": old_columns},
        )

        job = executed(validated(upload(payload, mode=spec.ImportMode.UPSERT)))

        assert job.status == ImportStatus.DONE
        variant = ProductVariant.objects.get(sku="GLV-M")
        assert variant.name_en == "New" and variant.price_adjustment == Decimal("3.00")
        assert variant.color is not None, "ملف قديم الأعمدة لا يفكّ ربط اللون"


class TestStockSheet:
    def stock_row(self, sku="GLV", variant_sku=None, quantity=5):
        row = {"sku": sku, "quantity": quantity, "unit_cost": "2.00"}
        if variant_sku:
            row["variant_sku"] = variant_sku
        return row

    def test_stock_goes_to_the_named_colour_variant(self, make_product_row, location):
        run(make_product_row, [variant_row("GLV-M")])
        payload = workbook(
            {
                "products": [make_product_row("GLV")],
                "stock": [self.stock_row(variant_sku="GLV-M")],
            }
        )

        job = executed(validated(upload(payload, mode=spec.ImportMode.UPSERT)))

        assert job.status == ImportStatus.DONE
        from inventory.models import Stock

        assert Stock.objects.get(variant__sku="GLV-M").quantity_physical == 5

    def test_stock_without_a_variant_is_refused_for_a_product_with_colours(
        self, make_product_row, location
    ):
        """⚠️  Units on "no variant" would be unsellable and nobody would notice."""
        run(make_product_row, [variant_row("GLV-M")])
        payload = workbook({"products": [make_product_row("GLV")], "stock": [self.stock_row()]})

        job = validated(upload(payload, mode=spec.ImportMode.UPSERT))

        assert job.status == ImportStatus.REJECTED
        assert job.errors.filter(message__contains="حدّد رمز النسخة").exists()

    def test_a_variant_of_another_product_is_refused(self, make_product_row, location):
        run(
            make_product_row,
            [variant_row("A-M", parent="A")],
            products=[make_product_row("A"), make_product_row("B")],
        )
        payload = workbook(
            {
                "products": [make_product_row("B")],
                "stock": [self.stock_row("B", "A-M")],
            }
        )

        job = validated(upload(payload, mode=spec.ImportMode.UPSERT))

        assert job.status == ImportStatus.REJECTED
        assert job.errors.filter(message__contains="منتجًا آخر").exists()

    def test_a_plain_product_still_takes_stock_without_a_variant(self, make_product_row, location):
        payload = workbook(
            {"products": [make_product_row("PLAIN")], "stock": [self.stock_row("PLAIN")]}
        )

        assert executed(validated(upload(payload))).status == ImportStatus.DONE


class TestTemplate:
    def test_the_colour_columns_are_in_the_spec_and_the_template(self):
        from io import BytesIO

        from openpyxl import load_workbook

        from imports import template

        keys = [column.key for column in spec.VARIANT_COLUMNS]
        assert keys[-4:] == list(spec.VARIANT_COLOR_KEYS)

        book = load_workbook(BytesIO(template.build()))
        headers = (
            [cell.value for cell in book[spec.SHEET_TITLES[spec.Sheet.VARIANTS]][1]]
            if (spec.SHEET_TITLES[spec.Sheet.VARIANTS] in book.sheetnames)
            else [cell.value for cell in book[spec.Sheet.VARIANTS][1]]
        )
        assert "لون العرض" in headers and "سعر اللون" in headers
