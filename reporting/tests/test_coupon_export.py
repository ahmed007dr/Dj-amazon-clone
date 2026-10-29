"""
The `coupons` dataset — the Excel version of the coupon analytics screen.

⚠️  Two things are pinned here that the promotions tests cannot reach, because
    `promotions` is below `orders` and `reporting` is the layer that may see both:

        · the status strings `promotions.analytics` compares against are the real
          `orders.models` values (a rename there must fail a test, not skew a report)
        · the sheet and the screen agree — same function, same numbers
"""

import io
from decimal import Decimal

import pytest
from django.apps import apps
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import AccountType, User
from catalog.models import Category, Product
from core.testing import grant_all_domains
from orders.models import OrderStatus, PaymentStatus
from promotions import analytics
from promotions.tests.conftest import make_coupon, make_order, make_user
from reporting.export import registry

pytestmark = pytest.mark.django_db

PASSWORD = "Str0ng-Test-Pass!23"


@pytest.fixture
def admin_client(db):
    admin = User.objects.create_user(
        email="coupon-export@test.local", password=PASSWORD, account_type=AccountType.ADMIN
    )
    admin.is_active = True
    admin.save()
    apps.get_model("administration", "AdminProfile").objects.create(user=admin)
    grant_all_domains(admin)

    client = APIClient()
    client.force_authenticate(user=admin)
    return client


@pytest.fixture
def product(db):
    category = Category.objects.create(slug="exp-cat", name_ar="فئة", name_en="Cat")
    return Product.objects.create(
        sku="EXP-1", name_ar="صنف", name_en="Item", category=category, base_price=Decimal("100")
    )


def download(client, **params):
    return client.get(reverse("v1:reporting:export-download", args=["coupons"]), params)


def rows_of(response):
    book = load_workbook(io.BytesIO(response.content))
    sheet = book[book.sheetnames[0]]
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    book.close()
    return values


def test_the_status_strings_match_the_orders_domain():
    """⚠️  `promotions` cannot import `orders`, so it holds copies — this keeps them honest."""
    assert analytics.STATUS_CANCELLED == OrderStatus.CANCELLED
    assert analytics.STATUS_REFUNDED == OrderStatus.REFUNDED
    assert analytics.PAYMENT_FAILED == PaymentStatus.FAILED


def test_the_dataset_is_registered_under_the_pricing_permission():
    from core.permissions import CanManagePricing

    dataset = registry.get("coupons")

    assert dataset is not None
    assert dataset.permission is CanManagePricing
    assert not dataset.has_contact


def test_every_row_matches_its_column_count(product):
    coupon = make_coupon("COLS", holder_name="Dr. Cols")
    make_order(
        coupon, make_user("cols@test.local"), [(product, 1, "100", "0")], coupon_discount="10.00"
    )
    dataset = registry.get("coupons")

    columns = dataset.columns(False)
    rows = list(dataset.rows({}, False))

    assert rows
    assert all(len(row) == len(columns) for row in rows)


def test_a_coupon_row_carries_holder_usage_and_sales(admin_client, product):
    coupon = make_coupon("DRX10", holder_name="Dr. Sara", holder_note="phone 0100000000")
    make_order(
        coupon,
        make_user("a@test.local"),
        [(product, 2, "100", "0")],
        coupon_discount="20.00",
        shipping="15.00",
    )
    make_order(
        coupon,
        make_user("b@test.local"),
        [(product, 1, "100", "0")],
        coupon_discount="10.00",
        status="CANCELLED",
        redemption_cancelled=True,
    )
    make_coupon("QUIET")

    response = download(admin_client)

    assert response.status_code == 200
    table = rows_of(response)
    headers_at = next(i for i, row in enumerate(table) if row and row[0] == "الكود")
    headers, body = table[headers_at], table[headers_at + 1 :]
    by_code = {row[0]: dict(zip(headers, row, strict=False)) for row in body if row and row[0]}

    row = by_code["DRX10"]
    assert row["الحامل"] == "Dr. Sara"
    assert row["مرات الاستخدام"] == 1
    assert row["الطلبات المحتسبة"] == 1
    assert row["العملاء"] == 1
    assert Decimal(str(row["المبيعات قبل الكوبون"])) == Decimal("200.00")
    assert Decimal(str(row["الخصم الممنوح"])) == Decimal("20.00")
    assert Decimal(str(row["صافي المبيعات"])) == Decimal("180.00")
    assert Decimal(str(row["إجمالي الطلبات"])) == Decimal("195.00")
    assert row["طلبات مستبعدة"] == 1
    assert row["الحالة"] == "جارٍ"

    assert by_code["QUIET"]["مرات الاستخدام"] == 0
    assert by_code["QUIET"]["الحامل"] in (None, "")


def test_the_sheet_agrees_with_the_screen(admin_client, product):
    """Same function, same number — the sheet is never a second implementation."""
    coupon = make_coupon("SAME")
    for index in range(3):
        make_order(
            coupon,
            make_user(f"same{index}@test.local"),
            [(product, index + 1, "33.33", "0")],
            coupon_discount="3.33",
        )
    screen = analytics.summaries([coupon.pk])[coupon.pk]

    dataset = registry.get("coupons")
    row = next(item for item in dataset.rows({}, False) if item[0] == "SAME")

    assert row[7] == screen["times_used"]
    assert row[8] == screen["orders"]
    assert row[10] == screen["gross_sales"]
    assert row[11] == screen["discount_given"]
    assert row[12] == screen["net_sales"]


def test_the_holder_note_is_never_exported(admin_client):
    make_coupon("NOTE", holder_name="X", holder_note="SECRET-FREE-TEXT-0100")

    response = download(admin_client)

    assert response.status_code == 200
    assert b"SECRET-FREE-TEXT-0100" not in response.content
    assert all("ملاحظة" not in column.header for column in registry.get("coupons").columns(False))


def test_the_search_filter_narrows_the_rows(admin_client):
    make_coupon("ONE", holder_name="Sara")
    make_coupon("TWO", holder_name="Omar")

    table = rows_of(download(admin_client, search="sara"))
    codes = {row[0] for row in table if row and row[0] in ("ONE", "TWO")}

    assert codes == {"ONE"}


def test_a_filter_with_no_rows_is_a_404_not_an_empty_file(admin_client):
    make_coupon("ONLY")
    assert download(admin_client, search="zzz-nothing").status_code == 404


def test_a_customer_cannot_download_it():
    client = APIClient()
    client.force_authenticate(user=make_user("plain@test.local"))

    assert download(client).status_code == 403


def test_it_is_not_even_listed_for_someone_without_the_permission():
    client = APIClient()
    client.force_authenticate(user=make_user("plain2@test.local"))

    response = client.get(reverse("v1:reporting:exports"))
    keys = {item["key"] for group in response.data.get("groups", []) for item in group["datasets"]}

    assert "coupons" not in keys


def test_a_bad_date_is_a_clear_400(admin_client):
    make_coupon("DATES")
    assert download(admin_client, start="2026-13-45").status_code == 400
    assert download(admin_client, start="2026-10-10", end="2026-10-01").status_code == 400


# ═══════════════════════════════════════════════════════════
#  Registered vs guest, new customers — the extra columns and the buyers sheet
# ═══════════════════════════════════════════════════════════


def table_of(response):
    """The first sheet as `(headers, [row dicts])`, found by the "الكود" header."""
    table = rows_of(response)
    at = next(i for i, row in enumerate(table) if row and row[0] == "الكود")
    headers = table[at]
    return headers, [dict(zip(headers, row, strict=False)) for row in table[at + 1 :] if row]


@pytest.fixture
def report_shop(product):
    from promotions.tests.test_coupon_report import guest_order, plain_order

    coupon = make_coupon("MIXED", holder_name="Dr. Mixed")
    new_user = make_user("export-new@test.local", name="جديد")
    old_user = make_user("export-old@test.local", name="قديم")
    plain_order(old_user, [(product, 1, "100", "0")], days_ago=9)

    make_order(coupon, new_user, [(product, 2, "100", "0")], coupon_discount="20.00")
    make_order(coupon, old_user, [(product, 1, "100", "0")], coupon_discount="10.00")
    guest_order(
        coupon, "export-guest@test.local", [(product, 1, "100", "0")], coupon_discount="10.00"
    )
    return coupon


def test_the_coupons_sheet_gains_the_buyer_columns_at_the_end(admin_client, report_shop):
    headers, body = table_of(download(admin_client))
    row = {r["الكود"]: r for r in body if r.get("الكود")}["MIXED"]

    assert headers[-4:] == [
        "طلبات المسجّلين",
        "طلبات الزوار",
        "عملاء جدد",
        "نسبة الخصم من المبيعات (%)",
    ]
    assert headers[0] == "الكود" and "الخصم الممنوح" in headers  # the old columns are untouched
    assert row["طلبات المسجّلين"] == 2
    assert row["طلبات الزوار"] == 1
    assert row["عملاء جدد"] == 2  # the new account and the first-time guest
    assert Decimal(str(row["نسبة الخصم من المبيعات (%)"])) == Decimal("10.00")  # 40 over 400


def test_the_sheet_agrees_with_the_screen_on_the_buyer_split(admin_client, report_shop):
    from promotions.analytics_customers import Breakdown

    block = Breakdown.load([report_shop.pk]).block([report_shop.pk])
    _headers, body = table_of(download(admin_client))
    row = {r["الكود"]: r for r in body if r.get("الكود")}["MIXED"]

    assert row["طلبات الزوار"] == block["by_customer_type"]["guest"]["orders"]
    assert row["عملاء جدد"] == block["new_customers"]["count"]


def test_the_buyers_dataset_is_registered_with_a_contact_gate():
    from core.permissions import CanManagePricing

    dataset = registry.get("coupon_customers")

    assert dataset is not None
    assert dataset.permission is CanManagePricing
    assert dataset.has_contact
    plain = [c.header for c in dataset.columns(False)]
    full = [c.header for c in dataset.columns(True)]
    assert full[: len(plain)] == plain
    assert {"الاسم", "البريد", "الهاتف"}.isdisjoint(plain)
    assert {"الاسم", "البريد", "الهاتف"} <= set(full)


def download_buyers(client, **params):
    return client.get(reverse("v1:reporting:export-download", args=["coupon_customers"]), params)


def test_the_buyers_sheet_has_one_row_per_coupon_and_buyer_without_contact(
    admin_client, report_shop
):
    response = download_buyers(admin_client)
    assert response.status_code == 200

    headers, body = table_of(response)
    rows = [r for r in body if r.get("الكود") == "MIXED"]

    assert len(rows) == 3
    assert "البريد" not in headers and "الاسم" not in headers
    kinds = sorted((r["نوع العميل"], r["عميل جديد"]) for r in rows)
    assert kinds == [("زائر", "نعم"), ("مسجّل", "لا"), ("مسجّل", "نعم")]
    assert sum(Decimal(str(r["الخصم الممنوح"])) for r in rows) == Decimal("40.00")
    assert b"export-guest@test.local" not in response.content


def test_the_buyers_sheet_with_contact_carries_names_and_emails(admin_client, report_shop):
    response = download_buyers(admin_client, include_contact="true")
    assert response.status_code == 200

    headers, body = table_of(response)
    emails = {r["البريد"] for r in body if r.get("الكود") == "MIXED"}

    assert {"الاسم", "البريد", "الهاتف"} <= set(headers)
    assert "export-guest@test.local" in emails
    assert "export-new@test.local" in emails


def test_contact_columns_need_the_accounts_permission(report_shop):
    """Pricing alone may read the buyers sheet, but not with names and e-mails."""
    from django.contrib.auth.models import Permission

    user = make_user("pricing-only@test.local")
    user.account_type = AccountType.ADMIN
    user.save()
    apps.get_model("administration", "AdminProfile").objects.create(user=user)
    user.user_permissions.add(Permission.objects.get(codename="change_pricelist"))
    client = APIClient()
    client.force_authenticate(user=User.objects.get(pk=user.pk))

    assert download_buyers(client).status_code == 200
    assert download_buyers(client, include_contact="true").status_code == 403


def test_the_buyers_sheet_filters_by_coupon_search_and_period(admin_client, report_shop, product):
    other = make_coupon("OTHER")
    make_order(other, make_user("o@test.local"), [(product, 1, "100", "0")], coupon_discount="5")

    _h, body = table_of(download_buyers(admin_client, search="MIXED"))
    assert {r["الكود"] for r in body if r.get("الكود")} == {"MIXED"}

    assert download_buyers(admin_client, start="2026-13-45").status_code == 400
    assert download_buyers(admin_client, start="2026-10-10", end="2026-10-01").status_code == 400
    assert download_buyers(admin_client, search="zzz-nothing").status_code == 404


def test_a_customer_cannot_download_the_buyers_sheet():
    client = APIClient()
    client.force_authenticate(user=make_user("plain3@test.local"))

    assert download_buyers(client).status_code == 403


# ═══════════════════════════════════════════════════════════
#  The referral members sheet
# ═══════════════════════════════════════════════════════════


@pytest.fixture
def referral_shop(db):
    from customers.models import CustomerProfile
    from loyalty.models import ReferralCode, ReferralProgram, ReferralStatus
    from loyalty.tests.test_referral_report import make_order as make_plain
    from loyalty.tests.test_referral_report import refer

    program = ReferralProgram.objects.create(
        code="exp-ref", name_ar="إحالة", name_en="Ref", is_active=True
    )
    member = make_user("export-member@test.local", name="منى")
    CustomerProfile.objects.get_or_create(user=member, defaults={"display_name_ar": "منى"})
    ReferralCode.objects.create(user=member, code="EXPM10")
    friend = make_user("export-friend@test.local", name="صديق")
    profile = CustomerProfile.objects.create(user=friend, display_name_ar="صديق")
    order = make_plain(profile, "250")
    refer(program, member, friend, ReferralStatus.REWARDED, order=order)
    pending = make_user("export-pending@test.local")
    CustomerProfile.objects.create(user=pending, display_name_ar="x")
    refer(program, member, pending, ReferralStatus.PENDING)
    return member


def download_members(client, **params):
    return client.get(reverse("v1:reporting:export-download", args=["referral_members"]), params)


def _member_row(response):
    table = rows_of(response)
    at = next(i for i, row in enumerate(table) if row and row[0] == "الترتيب")
    return table[at], dict(zip(table[at], table[at + 1], strict=False))


def test_the_referral_members_sheet(admin_client, referral_shop):
    response = download_members(admin_client)
    assert response.status_code == 200

    headers, row = _member_row(response)

    assert row["الترتيب"] == 1
    assert row["العضو"] == "منى"
    assert row["كود الإحالة"] == "EXPM10"
    assert row["المدعوون"] == 2
    assert row["المؤهَّلون"] == 1
    assert Decimal(str(row["نسبة التحويل (%)"])) == Decimal("50.00")
    assert Decimal(str(row["مبيعات العملاء الجدد"])) == Decimal("250.00")
    assert "البريد" not in headers
    assert b"export-member@test.local" not in response.content

    assert (
        b"export-member@test.local"
        in download_members(admin_client, include_contact="true").content
    )


def test_the_referral_sheet_agrees_with_the_screen(admin_client, referral_shop):
    from loyalty import report

    entry = report.member_stats()[referral_shop.pk]
    _headers, row = _member_row(download_members(admin_client))

    assert row["المدعوون"] == entry["invited"]
    assert row["النقاط الممنوحة"] == entry["points_awarded"]


def test_a_customer_cannot_download_the_referral_sheet():
    client = APIClient()
    client.force_authenticate(user=make_user("plain4@test.local"))

    assert download_members(client).status_code == 403
