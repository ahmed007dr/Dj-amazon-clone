"""
Export: permissions, secrets, the ceiling, the audit trail, and the round trip.

⚠️  The security tests here are the point of the file. An export is bulk
    disclosure in one request — the only feature in the panel where a single
    click can move the entire customer list out of the company.
"""

import io
from decimal import Decimal

import pytest
from django.apps import apps
from django.contrib.auth.models import Permission
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import AccountType, User
from catalog.models import Category, Product
from core.models.audit import AuditAction, AuditLog
from core.testing import grant_all_domains
from customers.models import CustomerProfile
from inventory.models import StockLocation
from reporting.export import registry, writer

pytestmark = pytest.mark.django_db

PASSWORD = "Str0ng-Test-Pass!23"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def make_admin(email: str, *, permissions: list[str] | None = None) -> APIClient:
    user = User.objects.create_user(email=email, password=PASSWORD, account_type=AccountType.ADMIN)
    user.is_active = True
    user.save()
    apps.get_model("administration", "AdminProfile").objects.create(user=user)

    if permissions is None:
        grant_all_domains(user)
        # ⚠️  `grant_all_domains` deliberately excludes finance — see its docstring.
        #     The finance datasets need it explicitly, and granting it here is
        #     what makes "the full admin" mean the same thing in this file as elsewhere.
        user.user_permissions.add(
            Permission.objects.get(
                content_type__app_label="finance", codename="view_revenueentry"
            )
        )
    else:
        for path in permissions:
            app_label, codename = path.split(".")
            user.user_permissions.add(
                Permission.objects.get(content_type__app_label=app_label, codename=codename)
            )

    client = APIClient()
    client.force_authenticate(user=user)
    return client


@pytest.fixture
def admin_client(db):
    return make_admin("exporter@test.local")


@pytest.fixture
def catalogue(db):
    category = Category.objects.create(name_ar="مستلزمات", name_en="Supplies")
    category.refresh_from_db()
    for n in range(3):
        Product.objects.create(
            sku=f"SKU-{n}",
            name_ar=f"صنف {n}",
            name_en=f"Item {n}",
            category=category,
            base_price=Decimal("12.50"),
        )
    return category


def download(client, key: str, **params):
    return client.get(reverse("v1:reporting:export-download", args=[key]), params)


def sheet_of(response, index: int = 0):
    book = load_workbook(io.BytesIO(response.content))
    return book, book[book.sheetnames[index]]


# ═══════════════════════════════════════════════════════════
#  The catalogue endpoint
# ═══════════════════════════════════════════════════════════


class TestCatalogue:
    def test_it_lists_the_datasets_grouped(self, admin_client, catalogue):
        body = admin_client.get(reverse("v1:reporting:exports")).json()

        keys = {
            dataset["key"] for group in body["groups"] for dataset in group["datasets"]
        }
        assert "stock-levels" in keys
        assert "customers" in keys
        assert body["max_rows"] == writer.MAX_EXPORT_ROWS

    def test_it_hides_what_the_user_may_not_export(self, db):
        """
        ⚠️  Filtered at listing time, not only at download. A screen offering a
            button that always 403s reads as a broken system — and the list of
            datasets is itself a disclosure about what the business tracks.
        """
        client = make_admin("warehouse@test.local", permissions=["inventory.change_stock"])
        body = client.get(reverse("v1:reporting:exports")).json()

        keys = {d["key"] for group in body["groups"] for d in group["datasets"]}
        assert "stock-levels" in keys
        assert "customers" not in keys
        assert "revenue" not in keys

    def test_the_contact_toggle_is_offered_only_to_who_could_use_it(self, db):
        client = make_admin("orders@test.local", permissions=["orders.change_order"])
        body = client.get(reverse("v1:reporting:exports")).json()

        orders = next(
            d for group in body["groups"] for d in group["datasets"] if d["key"] == "orders"
        )
        assert orders["contact_available"] is False

    def test_choice_filters_carry_their_values(self, admin_client, catalogue):
        StockLocation.objects.create(code="main", name_ar="المخزن", name_en="WH", is_default=True)
        body = admin_client.get(reverse("v1:reporting:exports")).json()

        codes = {row["value"] for row in body["options"]["locations"]}
        assert "main" in codes

    def test_an_anonymous_visitor_gets_nothing(self, db):
        assert APIClient().get(reverse("v1:reporting:exports")).status_code in (401, 403)


# ═══════════════════════════════════════════════════════════
#  The file
# ═══════════════════════════════════════════════════════════


class TestDownload:
    def test_it_returns_a_workbook(self, admin_client, catalogue):
        response = download(admin_client, "catalogue")

        assert response.status_code == 200
        assert response["Content-Type"] == XLSX
        assert response.content[:2] == b"PK"
        assert "attachment" in response["Content-Disposition"]

    def test_it_is_never_cached(self, admin_client, catalogue):
        """Built from live data for one authorised person."""
        assert download(admin_client, "catalogue")["Cache-Control"] == "no-store"

    def test_the_header_row_is_frozen_and_filterable(self, admin_client, catalogue):
        """
        ⚠️  In write-only mode a freeze assigned after the first row is silently
            dropped. This is the test that catches it — the file still opens, so
            nothing else would.
        """
        _, sheet = sheet_of(download(admin_client, "catalogue"))

        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None

    def test_money_is_a_number_and_not_text(self, admin_client, catalogue):
        """
        ⚠️  A text column answers no `SUM` and sorts alphabetically. The whole
            purpose of the file is analysis.
        """
        _, sheet = sheet_of(download(admin_client, "catalogue"))

        headers = [cell.value for cell in sheet[1]]
        cell = sheet[2][headers.index("السعر المرجعي")]

        assert isinstance(cell.value, int | float)
        assert cell.number_format == writer.MONEY_FORMAT

    def test_every_file_carries_its_provenance(self, admin_client, catalogue):
        book, _ = sheet_of(download(admin_client, "catalogue"))

        assert "_عن هذا الملف" in book.sheetnames
        meta = book["_عن هذا الملف"]
        labels = {row[0].value for row in meta.iter_rows(min_col=1, max_col=1)}
        assert "الفلاتر المطبَّقة" in labels
        assert "عدد الصفوف" in labels

    def test_filters_narrow_the_result(self, admin_client, catalogue):
        other = Category.objects.create(name_ar="أدوية", name_en="Medicines")
        other.refresh_from_db()
        Product.objects.create(
            sku="MED-1",
            name_ar="دواء",
            name_en="Medicine",
            category=other,
            base_price=Decimal("5.00"),
        )

        _, sheet = sheet_of(download(admin_client, "catalogue", category_path=other.path))
        assert sheet.max_row == 2  # header + the one product

    def test_no_matching_rows_is_a_message_not_an_empty_file(self, admin_client, catalogue):
        """An empty spreadsheet looks like a broken export rather than an empty result."""
        response = download(admin_client, "catalogue", category_path="no/such/path")
        assert response.status_code == 404

    def test_an_unknown_dataset_is_404(self, admin_client):
        assert download(admin_client, "invented").status_code == 404

    def test_a_bad_date_is_rejected_with_a_readable_message(self, admin_client, catalogue):
        response = download(admin_client, "orders", start="last-tuesday")
        assert response.status_code == 400
        assert "YYYY-MM-DD" in response.json()["detail"]

    def test_a_dataset_that_demands_a_period_refuses_without_one(self, admin_client, catalogue):
        """
        ⚠️  Stock movements and order lines grow without bound; an unfiltered
            export of them asks for the whole history of the business.
        """
        response = download(admin_client, "stock-movements")
        assert response.status_code == 400
        assert "فترة" in response.json()["detail"]

    def test_the_row_ceiling_names_the_real_count(self, admin_client, catalogue, monkeypatch):
        """
        ⚠️  "Too large" makes the admin guess how much to narrow. The number
            tells them.
        """
        monkeypatch.setattr(writer, "MAX_EXPORT_ROWS", 2)
        response = download(admin_client, "catalogue")

        assert response.status_code == 400
        assert "3" in response.json()["detail"]


# ═══════════════════════════════════════════════════════════
#  Permissions
# ═══════════════════════════════════════════════════════════


class TestPermissions:
    def test_a_warehouse_keeper_may_export_stock(self, db, catalogue):
        client = make_admin("wh@test.local", permissions=["inventory.change_stock"])
        StockLocation.objects.create(code="main", name_ar="المخزن", name_en="WH", is_default=True)

        # No stock rows yet — 404 for emptiness, not 403 for permission.
        assert download(client, "stock-levels").status_code == 404

    def test_a_warehouse_keeper_may_not_export_customers(self, db):
        client = make_admin("wh2@test.local", permissions=["inventory.change_stock"])
        response = download(client, "customers")

        assert response.status_code == 403
        assert "صلاحية" in response.json()["detail"]

    def test_the_catalogue_permission_does_not_open_finance(self, db, catalogue):
        client = make_admin("merch@test.local", permissions=["catalog.change_product"])

        assert download(client, "catalogue").status_code == 200
        assert download(client, "revenue").status_code == 403


# ═══════════════════════════════════════════════════════════
#  Contact details — the second gate
# ═══════════════════════════════════════════════════════════


@pytest.fixture
def customer(db):
    user = User.objects.create_user(
        email="buyer@test.local", password=PASSWORD, phone="01000000000"
    )
    user.is_active = True
    user.save()
    return CustomerProfile.objects.create(user=user, display_name_ar="مشترٍ")


class TestContactGate:
    def test_contact_columns_are_absent_by_default(self, admin_client, customer):
        """
        ⚠️  A customer export without contact details is analysis. The same file
            with them is a mailing list leaving the company.
        """
        _, sheet = sheet_of(download(admin_client, "customers"))
        headers = [cell.value for cell in sheet[1]]

        assert "البريد" not in headers
        assert "الهاتف" not in headers

    def test_they_appear_when_asked_for_with_the_permission(self, admin_client, customer):
        _, sheet = sheet_of(download(admin_client, "customers", include_contact="true"))
        headers = [cell.value for cell in sheet[1]]

        assert "البريد" in headers
        assert sheet[2][headers.index("البريد")].value == "buyer@test.local"

    def test_asking_without_the_permission_is_refused_not_stripped(self, db, customer):
        """
        ⚠️  Silently dropping the columns hands back a file the admin believes is
            complete. They mail it on, the recipient reports missing numbers, and
            nobody can tell whether the data or the export is at fault.
        """
        client = make_admin("orders2@test.local", permissions=["orders.change_order"])
        response = download(client, "orders", include_contact="true")

        assert response.status_code == 403
        assert "الحسابات" in response.json()["detail"]

    def test_asking_on_a_dataset_without_contact_columns_is_an_error(
        self, admin_client, catalogue
    ):
        response = download(admin_client, "catalogue", include_contact="true")
        assert response.status_code == 400

    def test_the_two_files_are_otherwise_column_for_column_identical(
        self, admin_client, customer
    ):
        """A pivot table built on one keeps working on the other."""
        _, plain = sheet_of(download(admin_client, "customers"))
        _, full = sheet_of(download(admin_client, "customers", include_contact="true"))

        plain_headers = [cell.value for cell in plain[1]]
        full_headers = [cell.value for cell in full[1]]

        assert full_headers[: len(plain_headers)] == plain_headers


# ═══════════════════════════════════════════════════════════
#  Secrets
# ═══════════════════════════════════════════════════════════


class TestSecrets:
    #: Substrings that must never appear in any header of any export.
    FORBIDDEN = ("password", "كلمة المرور", "token", "توكن", "api_key", "secret", "hash")

    def test_no_dataset_exposes_a_secret_in_its_headers(self, admin_client):
        """
        ⚠️  Swept across **every registered dataset**, including ones added later.

            A per-dataset test only guards the datasets somebody remembered to
            write a test for, and the failure this catches is precisely the one
            nobody remembers.
        """
        offenders = []
        for dataset in registry.all_datasets():
            for include_contact in (False, True):
                for column in dataset.columns(include_contact):
                    lowered = column.header.lower()
                    if any(word in lowered for word in self.FORBIDDEN):
                        offenders.append(f"{dataset.key}:{column.header}")

        assert not offenders, offenders

    def test_the_account_export_carries_no_password_field(self, admin_client, customer):
        response = download(admin_client, "accounts", include_contact="true")
        book, sheet = sheet_of(response)

        headers = [cell.value for cell in sheet[1]]
        assert "البريد" in headers  # it did include contact
        # and the hash is nowhere in the bytes at all
        assert b"pbkdf2" not in response.content
        assert b"argon2" not in response.content
        book.close()

    def test_document_metadata_carries_no_file_path(self, admin_client):
        dataset = registry.get("documents")
        headers = [column.header for column in dataset.columns(False)]

        assert not any("ملف" in header or "رابط" in header for header in headers)


# ═══════════════════════════════════════════════════════════
#  The audit trail
# ═══════════════════════════════════════════════════════════


class TestAudit:
    def test_every_export_is_recorded(self, admin_client, catalogue):
        """
        ⚠️  `AuditAction.EXPORT` existed unused since the beginning. This is what
            it is for: "who downloaded the customer list the week before they
            resigned?" has no answer anywhere else.
        """
        before = AuditLog.objects.filter(action=AuditAction.EXPORT).count()
        download(admin_client, "catalogue")

        entries = AuditLog.objects.filter(action=AuditAction.EXPORT)
        assert entries.count() == before + 1

        entry = entries.order_by("-created_at").first()
        assert entry.changes["dataset"] == "catalogue"
        assert entry.changes["rows"] == 3
        assert entry.changes["include_contact"] is False

    def test_the_record_names_the_filters_used(self, admin_client, catalogue):
        download(admin_client, "catalogue", category_path=catalogue.path)

        entry = AuditLog.objects.filter(action=AuditAction.EXPORT).order_by("-created_at").first()
        assert entry.changes["filters"]["category_path"] == catalogue.path

    def test_the_record_distinguishes_a_contact_export(self, admin_client, customer):
        download(admin_client, "customers", include_contact="true")

        entry = AuditLog.objects.filter(action=AuditAction.EXPORT).order_by("-created_at").first()
        assert entry.changes["include_contact"] is True

    def test_a_refused_export_writes_nothing(self, db):
        client = make_admin("wh3@test.local", permissions=["inventory.change_stock"])
        before = AuditLog.objects.filter(action=AuditAction.EXPORT).count()

        download(client, "customers")

        assert AuditLog.objects.filter(action=AuditAction.EXPORT).count() == before


# ═══════════════════════════════════════════════════════════
#  The round trip
# ═══════════════════════════════════════════════════════════


class TestRoundTrip:
    def test_the_products_export_carries_the_import_template_columns(
        self, admin_client, catalogue
    ):
        """
        ⚠️  This is the reason the feature exists in this shape. Updating four
            thousand prices needs a file holding their SKUs, and the importer
            could consume a spreadsheet nobody was able to produce.
        """
        from imports import spec

        _, sheet = sheet_of(download(admin_client, "products"))
        headers = [cell.value for cell in sheet[1]]

        assert headers == [column.header for column in spec.PRODUCT_COLUMNS]

    def test_references_come_out_as_codes_the_importer_accepts(self, admin_client, catalogue):
        """Not "Supplies" — `supplies`, the path the importer resolves."""
        _, sheet = sheet_of(download(admin_client, "products"))

        headers = [cell.value for cell in sheet[1]]
        value = sheet[2][headers.index("مسار الفئة *")].value
        assert value == catalogue.path

    def test_an_exported_file_imports_back(self, admin_client, catalogue):
        """
        ⚠️  The end-to-end promise, tested end to end: export → change a price →
            upload on the import screen → the catalogue moves.

            Every part of this has a unit test above and none of them would
            catch a header the importer cannot match.
        """
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook

        from imports import services, spec
        from imports.models import ImportStatus

        response = download(admin_client, "products")
        source = load_workbook(io.BytesIO(response.content))
        exported = source[source.sheetnames[0]]

        # Rebuild the workbook the way the admin would save it, with the key row
        # the importer matches on.
        book = Workbook(write_only=True)
        worksheet = book.create_sheet(spec.Sheet.PRODUCTS)
        worksheet.append([cell.value for cell in exported[1]])
        worksheet.append([column.key for column in spec.PRODUCT_COLUMNS])

        price_at = [column.key for column in spec.PRODUCT_COLUMNS].index("base_price")
        for row in exported.iter_rows(min_row=2, values_only=True):
            values = list(row)
            values[price_at] = "99.00"
            worksheet.append(values)

        buffer = io.BytesIO()
        book.save(buffer)
        book.close()
        source.close()

        job = services.create_job(
            SimpleUploadedFile("round-trip.xlsx", buffer.getvalue()),
            mode=spec.ImportMode.UPDATE_ONLY,
        )
        services.start_validation(job)
        for _ in range(50):
            if not job.is_running:
                break
            services.advance(job)

        assert job.status == ImportStatus.VALIDATED, list(
            job.errors.values_list("column", "message")[:5]
        )

        services.start_execution(job)
        for _ in range(50):
            if not job.is_running:
                break
            services.advance(job)

        assert job.status == ImportStatus.DONE
        assert Product.objects.filter(base_price=Decimal("99.00")).count() == 3
