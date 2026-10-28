"""
The "price before discount" column in the catalogue exports.

⚠️  Two datasets carry it and they are built differently: the round-trip `products` sheet
    reads its columns from `imports.spec` (so it follows by construction), while the
    readable `catalogue` sheet has a hand-written header list **and** a hand-written row
    tuple. The second pair is what drifts — a value one place off sits under the wrong
    heading in a file that still opens — so each cell is pinned to its header here.
"""

import io
from decimal import Decimal

import pytest
from django.apps import apps
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import AccountType, User
from catalog.models import Category, Product
from core.testing import grant_all_domains
from imports import spec

pytestmark = pytest.mark.django_db

PASSWORD = "Str0ng-Test-Pass!23"


@pytest.fixture
def admin_client(db):
    user = User.objects.create_user(
        email="cmp-exporter@test.local", password=PASSWORD, account_type=AccountType.ADMIN
    )
    user.is_active = True
    user.save()
    apps.get_model("administration", "AdminProfile").objects.create(user=user)
    grant_all_domains(user)

    client = APIClient()
    client.force_authenticate(user=user)
    return client


@pytest.fixture
def products(db):
    category = Category.objects.create(name_ar="مستلزمات", name_en="Supplies")
    category.refresh_from_db()
    Product.objects.create(
        sku="A-ON-SALE",
        name_ar="مخفض",
        name_en="On sale",
        category=category,
        base_price=Decimal("80.00"),
        compare_at_price=Decimal("100.00"),
        weight_grams=250,
        pack_size="10 قطع",
    )
    Product.objects.create(
        sku="B-PLAIN",
        name_ar="عادي",
        name_en="Plain",
        category=category,
        base_price=Decimal("50.00"),
    )
    return category


def rows_of(client, key):
    response = client.get(reverse("v1:reporting:export-download", args=[key]))
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    sheet = book[book.sheetnames[0]]
    headers = [cell.value for cell in sheet[1]]
    return headers, {
        row[0]: dict(zip(headers, row, strict=True))
        for row in sheet.iter_rows(min_row=2, values_only=True)
    }


class TestReadableCatalogue:
    def test_each_value_sits_under_its_own_header(self, admin_client, products):
        headers, rows = rows_of(admin_client, "catalogue")
        sale, plain = rows["A-ON-SALE"], rows["B-PLAIN"]

        assert "السعر قبل الخصم" in headers
        assert headers.index("السعر قبل الخصم") == headers.index("السعر المرجعي") + 1
        assert (sale["السعر المرجعي"], sale["السعر قبل الخصم"]) == (80, 100)
        assert (plain["السعر المرجعي"], plain["السعر قبل الخصم"]) == (50, None)
        # the columns after the new one did not shift
        assert sale["الوزن بالجرام"] == 250
        assert sale["حجم العبوة"] == "10 قطع"
        assert sale["مفعّل"] is not None

    def test_every_row_is_as_wide_as_the_header(self, admin_client, products):
        headers, rows = rows_of(admin_client, "catalogue")

        assert rows and all(len(row) == len(headers) for row in rows.values())


class TestRoundTripSheet:
    def test_it_carries_the_spec_column_with_the_value(self, admin_client, products):
        headers, rows = rows_of(admin_client, "products")

        assert headers == [column.header for column in spec.PRODUCT_COLUMNS]
        assert rows["A-ON-SALE"]["السعر قبل الخصم"] == 100
        assert rows["B-PLAIN"]["السعر قبل الخصم"] is None

    def test_the_cell_is_a_number_not_text(self, admin_client, products):
        _, rows = rows_of(admin_client, "products")

        assert isinstance(rows["A-ON-SALE"]["السعر قبل الخصم"], int | float)

    def test_an_unchanged_export_imports_back_with_the_old_price_intact(
        self, admin_client, products
    ):
        """⚠️  export → upload on the import screen must not trip the new rule or drop the value."""
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook

        from imports import services
        from imports.models import ImportStatus

        response = admin_client.get(reverse("v1:reporting:export-download", args=["products"]))
        source = load_workbook(io.BytesIO(response.content))
        exported = source[source.sheetnames[0]]

        book = Workbook(write_only=True)
        worksheet = book.create_sheet(spec.Sheet.PRODUCTS)
        worksheet.append([cell.value for cell in exported[1]])
        worksheet.append([column.key for column in spec.PRODUCT_COLUMNS])
        for row in exported.iter_rows(min_row=2, values_only=True):
            worksheet.append(list(row))
        buffer = io.BytesIO()
        book.save(buffer)

        job = services.create_job(
            SimpleUploadedFile("again.xlsx", buffer.getvalue()),
            mode=spec.ImportMode.UPDATE_ONLY,
        )
        for step in (services.start_validation, services.start_execution):
            step(job)
            for _ in range(50):
                if not job.is_running:
                    break
                services.advance(job)

        assert job.status == ImportStatus.DONE, list(job.errors.values_list("column", "message"))
        assert Product.objects.get(sku="A-ON-SALE").compare_at_price == Decimal("100.00")
        assert Product.objects.get(sku="B-PLAIN").compare_at_price is None
