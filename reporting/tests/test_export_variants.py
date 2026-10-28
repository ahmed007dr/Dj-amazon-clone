"""
The variants export — column for column the import sheet, colours included.

⚠️  The row used to be a hand-written tuple while the header came from `imports.spec`,
    so adding a column to the sheet shifted every value under the wrong heading. These
    tests derive both sides from the spec, so the next added column cannot do that.
"""

import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from catalog import colors as color_services
from catalog.models import Category, Product, ProductVariant
from imports import services, spec
from imports.models import ImportStatus
from reporting.tests.test_export import admin_client, download, sheet_of  # noqa: F401


@pytest.fixture
def catalogue(db):
    category = Category.objects.create(name_ar="ملابس", name_en="Apparel")
    category.refresh_from_db()
    product = Product.objects.create(
        sku="EXV-1", name_ar="قميص", name_en="Top", category=category, base_price=Decimal("100.00")
    )
    blue = color_services.create_color(
        product, name_ar="أزرق", name_en="Blue", hex="#1E40AF", price_override=Decimal("120.00")
    )
    ProductVariant.objects.create(
        product=product,
        sku="EXV-1-BLUE-L",
        name_ar="L",
        name_en="L",
        color=blue,
        attributes={"size": "L"},
        price_adjustment=Decimal("5.00"),
    )
    ProductVariant.objects.create(
        product=product,
        sku="EXV-1-PLAIN",
        name_ar="عادي",
        name_en="Plain",
        attributes={"size": "M"},
    )
    return product


def rows_by_sku(response):
    _, sheet = sheet_of(response)
    headers = [cell.value for cell in sheet[1]]
    sku_at = [column.key for column in spec.VARIANT_COLUMNS].index("sku")
    return headers, {row[sku_at]: row for row in sheet.iter_rows(min_row=2, values_only=True)}


@pytest.mark.django_db
class TestColumns:
    def test_the_headers_are_the_import_sheets(self, admin_client, catalogue):  # noqa: F811
        headers, _rows = rows_by_sku(download(admin_client, "variants"))

        assert headers == [column.header for column in spec.VARIANT_COLUMNS]

    def test_every_row_is_as_wide_as_the_header(self, admin_client, catalogue):  # noqa: F811
        headers, rows = rows_by_sku(download(admin_client, "variants"))

        assert rows and all(len(row) == len(headers) for row in rows.values())

    def test_the_colour_cells_hold_the_colour(self, admin_client, catalogue):  # noqa: F811
        _headers, rows = rows_by_sku(download(admin_client, "variants", include_inactive="true"))
        keys = [column.key for column in spec.VARIANT_COLUMNS]
        blue = dict(zip(keys, rows["EXV-1-BLUE-L"], strict=True))
        plain = dict(zip(keys, rows["EXV-1-PLAIN"], strict=True))

        assert (blue["color_name_ar"], blue["color_name_en"], blue["color_hex"]) == (
            "أزرق",
            "Blue",
            "#1E40AF",
        )
        assert blue["price_override"] == 120
        assert blue["attributes"] == "size=L"
        assert blue["parent_sku"] == "EXV-1"
        assert blue["price_adjustment"] == 5
        assert not plain["color_name_en"] and not plain["color_hex"]
        assert plain["price_override"] is None


@pytest.mark.django_db
class TestRoundTrip:
    def test_an_exported_file_imports_back_and_a_price_edit_lands_on_the_colour(
        self,
        admin_client,  # noqa: F811
        catalogue,
    ):
        response = download(admin_client, "variants")
        book_in = load_workbook(io.BytesIO(response.content))
        exported = book_in[book_in.sheetnames[0]]
        keys = [column.key for column in spec.VARIANT_COLUMNS]

        book = Workbook(write_only=True)
        # The importer demands a products sheet in every file (`REQUIRED_SHEETS`), so the
        # products export travels alongside the variants one, untouched.
        products_in = load_workbook(io.BytesIO(download(admin_client, "products").content))
        products = products_in[products_in.sheetnames[0]]
        products_sheet = book.create_sheet(spec.Sheet.PRODUCTS)
        products_sheet.append([cell.value for cell in products[1]])
        products_sheet.append([column.key for column in spec.PRODUCT_COLUMNS])
        for row in products.iter_rows(min_row=2, values_only=True):
            products_sheet.append(list(row))

        worksheet = book.create_sheet(spec.Sheet.VARIANTS)
        worksheet.append([cell.value for cell in exported[1]])
        worksheet.append(keys)
        price_at = keys.index("price_override")
        for row in exported.iter_rows(min_row=2, values_only=True):
            values = list(row)
            if values[price_at] is not None:
                values[price_at] = "150.00"
            worksheet.append(values)
        buffer = io.BytesIO()
        book.save(buffer)
        book.close()

        job = services.create_job(
            SimpleUploadedFile("variants.xlsx", buffer.getvalue()),
            mode=spec.ImportMode.UPDATE_ONLY,
        )
        for step in (services.start_validation, services.start_execution):
            step(job)
            for _ in range(50):
                if not job.is_running:
                    break
                services.advance(job)
            assert job.status in (ImportStatus.VALIDATED, ImportStatus.DONE), list(
                job.errors.values_list("column", "message")[:5]
            )

        assert catalogue.colors.get().price_override == Decimal("150.00")
        assert catalogue.colors.count() == 1, "no second colour after the re-upload"
