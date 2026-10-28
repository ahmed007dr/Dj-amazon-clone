"""
Guest orders in reports, exports and targets.

⚠️  `values("customer").distinct().count()` counts the NULL group as ONE
    customer. Every counter sale and every guest order — however many, from
    however many strangers — used to add exactly 1 to "customers".

    The unit now is a buyer we can tell apart: an account, or a guest by email
    (case-insensitively). A counter sale identifies nobody and is not counted.
"""

from __future__ import annotations

import io
from datetime import timedelta
from decimal import Decimal

import pytest
from django.apps import apps
from django.contrib.auth.models import Permission
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import AccountType, User
from core.testing import grant_all_domains
from customers.models import CustomerProfile
from employees.models import EmployeeProfile, EmployeeRole, EmployeeRoleKind
from orders.models import Order, OrderChannel, OrderStatus, PaymentStatus
from reporting import services
from targets import services as target_services
from targets.models import MonthlyTarget, TargetStatus, TargetType

PASSWORD = "Str0ng-Test-Pass!23"


def make_order(*, customer=None, channel=OrderChannel.ONLINE, email="", total="100.00", **extra):
    amount = Decimal(total)
    return Order.objects.create(
        customer=customer,
        channel=channel,
        contact_email=email,
        status=extra.pop("status", OrderStatus.DELIVERED),
        payment_status=PaymentStatus.PAID,
        subtotal=amount,
        grand_total=amount,
        recipient_name=extra.pop("recipient_name", "مشترٍ"),
        recipient_phone=extra.pop("recipient_phone", ""),
        **extra,
    )


def make_account(email):
    user = User.objects.create_user(email=email, password=PASSWORD)
    user.is_active = True
    user.save()
    return CustomerProfile.objects.create(user=user, display_name_ar="حساب")


def window():
    today = timezone.localdate()
    return today - timedelta(days=1), today + timedelta(days=1)


# ═══════════════════════════════════════════════════════════
#  Counting buyers
# ═══════════════════════════════════════════════════════════


@pytest.mark.django_db
class TestSalesSummaryCustomers:
    def test_counter_sales_are_not_a_customer(self):
        for _ in range(3):
            make_order(channel=OrderChannel.POS)

        summary = services.sales_summary(*window())

        assert summary.orders_count == 3
        assert summary.customers_count == 0, "three anonymous sales are zero identified buyers"

    def test_each_guest_email_is_one_buyer_whatever_its_case(self):
        make_order(email="a@example.com")
        make_order(email="A@Example.com")  # the same person, typed differently
        make_order(email="b@example.com")

        assert services.sales_summary(*window()).customers_count == 2

    def test_accounts_and_guests_are_added(self):
        account = make_account("member@test.local")
        make_order(customer=account)
        make_order(customer=account)  # one account, two orders
        make_order(email="guest@example.com")
        make_order(channel=OrderChannel.POS)

        assert services.sales_summary(*window()).customers_count == 2

    def test_a_cancelled_guest_order_is_not_a_sale(self):
        make_order(email="gone@example.com", status=OrderStatus.CANCELLED)
        assert services.sales_summary(*window()).customers_count == 0

    def test_a_claimed_order_counts_once_as_the_account(self):
        account = make_account("claimed@example.com")
        make_order(customer=account, email="claimed@example.com")  # guest-origin, now owned
        make_order(customer=account)

        assert services.sales_summary(*window()).customers_count == 1


@pytest.mark.django_db
class TestCustomerBehaviour:
    def test_returning_buyers_include_guests_by_email(self):
        account = make_account("member@test.local")
        make_order(customer=account)
        make_order(customer=account)
        make_order(email="twice@example.com")
        make_order(email="TWICE@example.com")
        make_order(email="once@example.com")
        make_order(channel=OrderChannel.POS)
        make_order(channel=OrderChannel.POS)  # two anonymous sales are not a returning buyer

        result = services.customer_behaviour(*window())

        assert result["returning_customers"] == 2
        assert result["active_customers"] == 3  # the account, twice@, once@
        assert result["guest_customers"] == 2

    def test_the_account_only_lists_are_untouched(self):
        make_order(email="guest@example.com")
        result = services.customer_behaviour(*window())
        assert result["top_customers"] == []


# ═══════════════════════════════════════════════════════════
#  Targets
# ═══════════════════════════════════════════════════════════


@pytest.fixture
def rep(db):
    role = EmployeeRole.objects.create(
        code="rep", kind=EmployeeRoleKind.SALES_REP, name_ar="مندوب", name_en="Rep"
    )
    user = User.objects.create_user(
        email="rep@test.local", password=PASSWORD, account_type=AccountType.EMPLOYEE
    )
    user.is_active = True
    user.save()
    return EmployeeProfile.objects.create(user=user, employee_number="EMP-1", role=role)


@pytest.mark.django_db
def test_a_customer_count_target_is_not_met_by_anonymous_sales(rep):
    today = timezone.localdate()
    target = MonthlyTarget.objects.create(
        employee=rep,
        year=today.year,
        month=today.month,
        target_type=TargetType.CUSTOMER_COUNT,
        target_value=Decimal("5"),
        status=TargetStatus.ACTIVE,
    )
    account = make_account("buyer@test.local")
    make_order(customer=account, channel=OrderChannel.EMPLOYEE, owner_employee=rep.user)
    for _ in range(4):  # four counter sales the rep rang up: nobody identified
        make_order(channel=OrderChannel.POS, owner_employee=rep.user)

    achievement = target_services.measure(target)

    assert achievement.orders_count == 5
    assert achievement.customers_count == 1, "the NULL group used to count as a second customer"


# ═══════════════════════════════════════════════════════════
#  The export
# ═══════════════════════════════════════════════════════════


def make_admin(permissions=None, email="exporter@test.local"):
    user = User.objects.create_user(email=email, password=PASSWORD, account_type=AccountType.ADMIN)
    user.is_active = True
    user.save()
    apps.get_model("administration", "AdminProfile").objects.create(user=user)
    if permissions is None:
        grant_all_domains(user)
    else:
        for path in permissions:
            app_label, codename = path.split(".")
            user.user_permissions.add(
                Permission.objects.get(content_type__app_label=app_label, codename=codename)
            )
    client = APIClient()
    client.force_authenticate(user=user)
    return client


def download_orders(client, **params):
    response = client.get(reverse("v1:reporting:export-download", args=["orders"]), params)
    assert response.status_code == 200, response.content[:200]
    book = load_workbook(io.BytesIO(response.content))
    sheet = book[book.sheetnames[0]]
    headers = [cell.value for cell in sheet[1]]
    rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]
    return headers, rows


@pytest.mark.django_db
class TestOrdersExport:
    @pytest.fixture
    def orders(self):
        account = make_account("member@test.local")
        return {
            "member": make_order(customer=account, recipient_phone="+201155556666"),
            "guest": make_order(
                email="guest@example.com",
                recipient_name="زائر",
                recipient_phone="+201001234567",
            ),
            "counter": make_order(channel=OrderChannel.POS),
        }

    def test_the_buyer_type_column_is_always_there(self, orders):
        headers, rows = download_orders(make_admin())

        kind = headers.index("نوع المشتري")
        number = headers.index("رقم الطلب")
        by_number = {row[number]: row[kind] for row in rows}
        assert by_number[orders["member"].number] == "حساب"
        assert by_number[orders["guest"].number] == "زائر"
        assert by_number[orders["counter"].number] == "بلا هوية"

    def test_the_guest_email_is_absent_without_the_contact_flag(self, orders):
        headers, rows = download_orders(make_admin())

        assert "بريد الزائر" not in headers
        assert "guest@example.com" not in str(rows)
        assert "+201001234567" not in str(rows)

    def test_the_guest_email_appears_behind_the_same_gate_as_the_phone(self, orders):
        headers, rows = download_orders(make_admin(), include_contact="true")

        assert "بريد الزائر" in headers
        email = headers.index("بريد الزائر")
        number = headers.index("رقم الطلب")
        by_number = {row[number]: row[email] for row in rows}
        assert by_number[orders["guest"].number] == "guest@example.com"
        assert not by_number[orders["member"].number]

    def test_asking_for_contact_without_the_permission_is_refused(self, orders):
        client = make_admin(["orders.change_order"])
        response = client.get(
            reverse("v1:reporting:export-download", args=["orders"]), {"include_contact": "true"}
        )
        assert response.status_code == 403

    def test_the_plain_file_is_a_prefix_of_the_full_one(self, orders):
        """A pivot built on one keeps working on the other."""
        plain, _ = download_orders(make_admin())
        full, _ = download_orders(make_admin(email="second@test.local"), include_contact="true")
        assert full[: len(plain)] == plain
